#!/usr/bin/env python3
"""
PAL/NTSC Video IP Core — Change Log & Working Features Note
Generates CHANGE_LOG.docx
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = '/home/user/video-ip-core-/tools/CHANGE_LOG.docx'

# ── colours ──────────────────────────────────────────────────────────────────
C_TITLE   = RGBColor(0x1F, 0x49, 0x7D)
C_HDR     = RGBColor(0x1F, 0x49, 0x7D)
C_HDR2    = RGBColor(0x2E, 0x74, 0xB5)
C_ALT     = RGBColor(0xDA, 0xE8, 0xFC)
C_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
C_BLACK   = RGBColor(0x00, 0x00, 0x00)
C_GREEN   = RGBColor(0x37, 0x86, 0x10)
C_RED     = RGBColor(0xC0, 0x00, 0x00)
C_ORANGE  = RGBColor(0xC5, 0x5A, 0x11)
C_BLUE    = RGBColor(0x2E, 0x74, 0xB5)
C_PURPLE  = RGBColor(0x70, 0x30, 0xA0)

# ── helpers ───────────────────────────────────────────────────────────────────
def set_cell_bg(cell, rgb):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  f'{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}')
    tcPr.append(shd)

def cp(cell, text, bold=False, color=C_BLACK, size=9.5, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.alignment = align
    r = p.add_run(text)
    r.bold = bold; r.italic = italic
    r.font.size = Pt(size)
    r.font.color.rgb = color

def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f'Heading {level}'
    r = p.add_run(text)
    r.font.color.rgb = C_TITLE if level == 1 else C_HDR2

def body(doc, text, size=10, color=C_BLACK, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.bold = bold; r.italic = italic
    return p

def status_badge(cell, text, color):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f'  {text}  ')
    r.bold = True; r.font.size = Pt(9)
    r.font.color.rgb = C_WHITE
    set_cell_bg(cell, color)

def build_table(doc, headers, rows, col_widths, zebra=True):
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.style = 'Table Grid'
    hrow = t.rows[0]
    for i,h in enumerate(headers):
        set_cell_bg(hrow.cells[i], C_HDR)
        cp(hrow.cells[i], h, bold=True, color=C_WHITE, size=10)
    for ri, row in enumerate(rows):
        bg = C_ALT if zebra and ri%2==1 else C_WHITE
        for ci, val in enumerate(row):
            set_cell_bg(t.rows[ri+1].cells[ci], bg)
            if isinstance(val, tuple):
                txt, col, bld = val
                cp(t.rows[ri+1].cells[ci], txt, bold=bld, color=col, size=9.5)
            else:
                cp(t.rows[ri+1].cells[ci], str(val), size=9.5)
    for ri in range(len(t.rows)):
        for ci,w in enumerate(col_widths):
            t.rows[ri].cells[ci].width = Inches(w)
    doc.add_paragraph()
    return t

# ═════════════════════════════════════════════════════════════════════════════
doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2)
    s.left_margin = s.right_margin = Cm(2.2)

# ── Title ────────────────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('PAL / NTSC Video IP Core')
r.bold = True; r.font.size = Pt(22); r.font.color.rgb = C_TITLE

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Change Log & Working Features')
r.bold = True; r.font.size = Pt(16); r.font.color.rgb = C_HDR2

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Branch: claude/xilinx-pal-video-ip-RIDhJ     Last updated: 2026-05-24')
r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x60,0x60,0x60)

doc.add_page_break()

# ═════════════════════════════════════════════════════════════════════════════
# 1. FEATURE STATUS
# ═════════════════════════════════════════════════════════════════════════════
heading(doc, '1  Feature Status — All Working', 1)

FEATURES = [
    # (Feature, File(s) changed, Status, Notes)
    ('sel 0 — Vertical stripes',
     'pal_tv_bram_top.vhd\nvideo_bram_top.vhd',
     'WORKING', '8-zone stripe pattern; zone width = STRIPE_W generic'),

    ('sel 1 — Horizontal stripes',
     'pal_tv_bram_top.vhd\nvideo_bram_top.vhd',
     'WORKING', '8-zone stripe pattern; zone height adapts PAL/NTSC'),

    ('sel 2 — Horizontal gradient',
     'pal_tv_bram_top.vhd\nvideo_bram_top.vhd',
     'WORKING', '10 luminance steps left→right'),

    ('sel 3 — Vertical gradient',
     'pal_tv_bram_top.vhd\nvideo_bram_top.vhd',
     'WORKING', '10 luminance steps top→bottom'),

    ('sel 4 — BRAM image (single/double buffer)',
     'pal_tv_bram_top.vhd\nvideo_bram_top.vhd',
     'WORKING', 'Natural sequential row order; interlace-corrected read; double-buffer added'),

    ('sel 5 — Bouncing ball',
     'pal_tv_bram_top.vhd\nvideo_bram_top.vhd',
     'WORKING', 'Runtime size/speed via ball_w_i, ball_h_i, ball_spd_h, ball_spd_v ports'),

    ('sel 6 — Full white',
     'pal_tv_bram_top.vhd\nvideo_bram_top.vhd',
     'WORKING', 'Every active pixel = LEVEL_WHITE; use for DAC peak-white calibration'),

    ('sel 7 — Full black',
     'pal_tv_bram_top.vhd\nvideo_bram_top.vhd',
     'WORKING', 'Every active pixel = LEVEL_BLANK; use for DAC black-level calibration'),

    ('sel 8 — Crosshatch grid',
     'pal_tv_bram_top.vhd\nvideo_bram_top.vhd',
     'WORKING', 'Runtime H/V pitch via cross_h_i / cross_v_i; modulo counter (no divider)'),

    ('sel 9 — Centre cross',
     'pal_tv_bram_top.vhd\nvideo_bram_top.vhd',
     'WORKING', 'Single white vertical line at H_ACTIVE/2; single white horizontal at V centre'),

    ('Ball control ports',
     'pal_tv_bram_top.vhd\nvideo_bram_top.vhd',
     'WORKING', 'ball_w_i, ball_h_i (10-bit); ball_spd_h, ball_spd_v (4-bit). Latch with rst pulse'),

    ('Crosshatch control ports',
     'pal_tv_bram_top.vhd\nvideo_bram_top.vhd',
     'WORKING', 'cross_h_i, cross_v_i (10-bit). Change any time; min 1 enforced internally'),

    ('BRAM double-buffering',
     'pal_tv_bram_top.vhd\nvideo_bram_top.vhd',
     'WORKING', 'buf_swap pulse → swap at next V-blank. buf_swapped strobe confirms. back_buf_o shows write bank'),

    ('PAL/NTSC runtime select',
     'video_bram_top.vhd',
     'WORKING', 'ntsc_mode port: 0=PAL 625/50, 1=NTSC 525/60. All timing adapts at runtime'),

    ('brightness / black_lvl ports',
     'all top files',
     'WORKING', '4-bit DAC level override. Clamped: black_lvl ≥ 4, black_lvl ≤ brightness'),

    ('bram_len port',
     'pal_tv_bram_top.vhd\nvideo_bram_top.vhd',
     'WORKING', '19-bit. Sets read address wrap point. 520=1 row loop, 299520=full frame'),

    ('line_sync_o / frame_sync_o outputs',
     'all bram/interlaced files',
     'WORKING', 'Separate H-sync and V-sync outputs alongside composite csync_o'),

    ('field_o output',
     'all bram/interlaced files',
     'WORKING', '0 = Field 1 (odd lines), 1 = Field 2 (even lines)'),

    ('HTML pattern generator tool',
     'tools/bram_pattern_gen.html',
     'WORKING', 'Browser-based. Patterns, text overlay, multi-zone, PNG export, 1-bit binary BRAM export'),

    ('User & Reference Guide (Word)',
     'tools/PAL_NTSC_Video_IP_User_Guide.docx',
     'WORKING', 'Full 10-section guide with simulation logs, waveform images, DAC strip diagrams'),

    ('Configuration Summary (Word)',
     'tools/PAL_NTSC_Config_Summary.docx',
     'WORKING', 'All generics, runtime ports, sel table, output ports, quick-reference'),
]

t = doc.add_table(rows=1+len(FEATURES), cols=4)
t.style = 'Table Grid'
for i,h in enumerate(['Feature', 'File(s)', 'Status', 'Notes']):
    set_cell_bg(t.rows[0].cells[i], C_HDR)
    cp(t.rows[0].cells[i], h, bold=True, color=C_WHITE, size=10)

for ri, (feat, files, status, notes) in enumerate(FEATURES):
    bg = C_ALT if ri%2==1 else C_WHITE
    for ci in range(4): set_cell_bg(t.rows[ri+1].cells[ci], bg)
    cp(t.rows[ri+1].cells[0], feat,   bold=True,  color=C_TITLE,  size=9.5)
    cp(t.rows[ri+1].cells[1], files,  bold=False, color=C_BLACK,  size=8.5, italic=True)
    cp(t.rows[ri+1].cells[2], status, bold=True,  color=C_GREEN,  size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    cp(t.rows[ri+1].cells[3], notes,  bold=False, color=C_BLACK,  size=9.5)

for ri in range(len(t.rows)):
    for ci,w in enumerate([2.0, 1.5, 0.8, 3.0]):
        t.rows[ri].cells[ci].width = Inches(w)
doc.add_paragraph()

doc.add_page_break()

# ═════════════════════════════════════════════════════════════════════════════
# 2. CHANGE HISTORY
# ═════════════════════════════════════════════════════════════════════════════
heading(doc, '2  Change History', 1)

CHANGES = [
    # (Change, What was done, Files touched)
    ('BRAM double-buffering\n(latest)',
     '• Two BRAM banks: bram_mem_0 and bram_mem_1\n'
     '• Write port always targets back buffer (NOT displayed)\n'
     '• buf_swap pulse deferred to V-blank boundary — no tearing\n'
     '• buf_swapped strobe confirms swap completion\n'
     '• back_buf_o output tells host which bank to write to\n'
     '• buf_swap must be a ONE-clock pulse; holding high re-arms every frame\n'
     '• Testbench: tb_bram_double_buf.vhd — 4 phases, all PASS',
     'pal_tv_bram_top.vhd\nvideo_bram_top.vhd\ntb_bram_double_buf.vhd'),

    ('Configuration summary document',
     '• PAL_NTSC_Config_Summary.docx: all generics, runtime ports,\n'
     '  sel table, output ports, quick-reference summary\n'
     '• Generator: tools/gen_config_summary.py',
     'tools/PAL_NTSC_Config_Summary.docx\ntools/gen_config_summary.py'),

    ('User Guide v2 — simulation results',
     '• Waveform PNG images from GHDL VCD files via matplotlib\n'
     '• DAC output strip diagrams for all sel values\n'
     '• Simulation console logs (PASS/FAIL colour coded)\n'
     '• Generator: tools/gen_report_v2.py',
     'tools/PAL_NTSC_Video_IP_User_Guide.docx\ntools/gen_report_v2.py'),

    ('sel 6–9 + ball/crosshatch ports',
     '• sel 6: Full white (DAC peak-white calibration)\n'
     '• sel 7: Full black (DAC black-level calibration)\n'
     '• sel 8: Crosshatch grid with runtime H/V pitch ports (cross_h_i, cross_v_i)\n'
     '• sel 9: Centre cross (H_ACTIVE/2 vertical + V_ACTIVE/2 horizontal)\n'
     '• Ball ports: ball_w_i, ball_h_i (width/height), ball_spd_h, ball_spd_v (speed)\n'
     '• Crosshatch uses modulo counters — no hardware division\n'
     '• Testbench: tb_new_patterns.vhd — 5 checks, all PASS',
     'pal_tv_bram_top.vhd\nvideo_bram_top.vhd\ntb_new_patterns.vhd'),

    ('HTML pattern generator tool v2',
     '• Text overlay: multiple entries, position, font size, colour, bold/italic\n'
     '• Multi-zone pattern: V/H split, per-zone type, colours, stripe width\n'
     '• Separate colours section with FG/BG pickers\n'
     '• PNG export and 1-bit binary BRAM export\n'
     '• Renders text before export (baked into image)',
     'tools/bram_pattern_gen.html'),

    ('User Guide v1 — base document',
     '• 12-section Word document\n'
     '• File load order, port tables, pattern descriptions\n'
     '• GHDL simulation commands, Vivado integration notes\n'
     '• Quick-start guide',
     'tools/PAL_NTSC_Video_IP_User_Guide.docx\ntools/gen_report.py'),

    ('PAL/NTSC runtime select (video_bram_top)',
     '• New top-level entity video_bram_top\n'
     '• ntsc_mode port: 0=PAL 625/50, 1=NTSC 525/60\n'
     '• All timing signals driven combinationally from ntsc_mode\n'
     '• H_TOTAL=640 PAL / 636 NTSC; V_TOTAL=625 PAL / 525 NTSC',
     'tv/bram/video_bram_top.vhd'),

    ('line_sync_o, frame_sync_o output ports',
     '• Separate H-sync and V-sync outputs added to all variants\n'
     '• line_sync_o: H pulse only, suppressed during vsync\n'
     '• frame_sync_o: HIGH for entire vsync region\n'
     '• field_o: 0=F1, 1=F2',
     'pal_tv_bram_top.vhd\nvideo_bram_top.vhd\npal_tv_interlaced_top.vhd'),

    ('sel 5 — Bouncing ball animation',
     '• White rectangle bouncing inside active picture area\n'
     '• Elastic bounce (velocity reflects on wall contact)\n'
     '• Default: 60×50 px, speed 3h/2v px per frame',
     'tv/bram/pal_tv_bram_top.vhd'),

    ('sel 4 — BRAM natural sequential row order',
     '• BRAM addressed in natural image order (row 0,1,2…)\n'
     '• Hardware routes to correct interlaced field automatically\n'
     '• bram_len port for partial fill (loop shorter than full frame)\n'
     '• pal_bram_loader.vhd: FIFO-to-BRAM loader utility',
     'tv/bram/pal_tv_bram_top.vhd\ntv/bram/pal_bram_loader.vhd'),

    ('brightness / black_lvl ports',
     '• Replace fixed DAC levels with runtime-adjustable ports\n'
     '• brightness: white DAC level (4-bit, clamped ≥ 4)\n'
     '• black_lvl: black/blanking floor (clamped 4 ≤ x ≤ brightness)',
     'all top-level files'),
]

for ci_idx, (title, details, files) in enumerate(CHANGES):
    # section box
    t2 = doc.add_table(rows=1, cols=3)
    t2.style = 'Table Grid'
    bg = C_HDR if ci_idx == 0 else C_HDR2
    set_cell_bg(t2.rows[0].cells[0], bg)
    cp(t2.rows[0].cells[0], title,   bold=True,  color=C_WHITE,  size=9.5)
    set_cell_bg(t2.rows[0].cells[1], C_WHITE if ci_idx%2==0 else C_ALT)
    cp(t2.rows[0].cells[1], details, bold=False, color=C_BLACK,  size=9.5)
    set_cell_bg(t2.rows[0].cells[2], C_WHITE if ci_idx%2==0 else C_ALT)
    cp(t2.rows[0].cells[2], files,   bold=False, color=C_BLUE,   size=8.5, italic=True)
    for ci,w in enumerate([1.5, 4.2, 1.6]):
        t2.rows[0].cells[ci].width = Inches(w)
    doc.add_paragraph()

doc.add_page_break()

# ═════════════════════════════════════════════════════════════════════════════
# 3. DOUBLE-BUFFER QUICK REFERENCE
# ═════════════════════════════════════════════════════════════════════════════
heading(doc, '3  Double-Buffer Quick Reference', 1)

body(doc, 'Key rules to remember:', bold=True, size=10)
doc.add_paragraph()

RULES = [
    ('Buffer does NOT auto-swap',
     'The display reads the same bank forever unless you explicitly pulse buf_swap.'),
    ('Pulse buf_swap for ONE clock only',
     'Set buf_swap=1 for one clock then immediately return to 0.\n'
     'Holding it high re-arms after each swap — causes a new swap every frame.'),
    ('Swap executes at V-blank, not immediately',
     'No matter when you pulse buf_swap (start/middle/end of frame),\n'
     'the actual bank flip happens at the last pixel of the last F2 active line.\n'
     'This guarantees zero tearing.'),
    ('Wait for buf_swapped before writing next frame',
     'buf_swapped strobes 1 for one clock when swap completes.\n'
     'Only after this strobe is it safe to start writing the next image.'),
    ('Always check back_buf_o before writing',
     'back_buf_o tells you which bank is currently the back (write) buffer.\n'
     'Always write to this bank. It automatically changes after each swap.'),
    ('Multiple pulses before V-blank = one swap',
     'Pulsing buf_swap twice before V-blank is fine — swap_req is idempotent.\n'
     'Only one swap fires per V-blank period.'),
]

build_table(doc,
    headers=['Rule', 'Detail'],
    rows=[(f'★  {r}', d) for r, d in RULES],
    col_widths=[2.2, 5.1])

# ═════════════════════════════════════════════════════════════════════════════
# 4. SIMULATION RESULTS SUMMARY
# ═════════════════════════════════════════════════════════════════════════════
heading(doc, '4  Simulation Results Summary', 1)

SIM = [
    ('tb_new_patterns',      'tb_new_patterns.vhd',      '5/5', 'sel=6 white, sel=7 black, sel=8 crosshatch, sel=9 centre cross, sel=5 ball ports'),
    ('tb_bram_double_buf',   'tb_bram_double_buf.vhd',   '4/4', 'Back buf write, swap delivery, alternating pattern, one-swap-per-V-blank'),
    ('tb_pal_tv_bram_top',   'tb_pal_tv_bram_top.vhd',   'PASS', 'BRAM interlaced addressing, 520×576 full frame'),
    ('tb_pal_sel5',          'tb_pal_sel5.vhd',           'PASS', 'Bouncing ball initial position and size'),
    ('tb_pal_bram_loader',   'tb_pal_bram_loader.vhd',   'PASS', 'BRAM loader FIFO-to-BRAM transfer'),
    ('tb_pal_tv_interlaced', 'tb_pal_tv_interlaced_top.vhd','PASS', 'PAL 625/50 timing, csync, field, active'),
    ('tb_pal_tv_crt50',      'tb_pal_tv_crt50_top.vhd',  'PASS', 'CRT50 progressive timing'),
    ('tb_video_bram_top',    'tb_video_bram_top.vhd',    'PASS', 'PAL/NTSC runtime switch, timing constants'),
]

build_table(doc,
    headers=['Testbench', 'File', 'Result', 'Coverage'],
    rows=[(n, f, (r, C_GREEN, True), c) for n,f,r,c in SIM],
    col_widths=[1.6, 1.8, 0.7, 3.2])

# ── save ─────────────────────────────────────────────────────────────────────
doc.save(OUT)
import os
print(f'Saved: {OUT}  ({os.path.getsize(OUT):,} bytes)')
