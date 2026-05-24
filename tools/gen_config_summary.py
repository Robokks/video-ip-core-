#!/usr/bin/env python3
"""
PAL/NTSC Video IP Core – Configuration Summary Document
Generates a clean Word .docx with all generics, ports, sel values and outputs.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = '/home/user/video-ip-core-/tools/PAL_NTSC_Config_Summary.docx'

# ─── colour palette ──────────────────────────────────────────────────────────
C_TITLE    = RGBColor(0x1F, 0x49, 0x7D)   # dark blue
C_HDR_BG   = RGBColor(0x1F, 0x49, 0x7D)   # table header bg
C_HDR2_BG  = RGBColor(0x2E, 0x74, 0xB5)   # sub-header bg
C_ROW_ALT  = RGBColor(0xDA, 0xE8, 0xFC)   # alternating row
C_WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
C_BLACK    = RGBColor(0x00, 0x00, 0x00)
C_GREEN    = RGBColor(0x37, 0x86, 0x10)
C_ORANGE   = RGBColor(0xC5, 0x5A, 0x11)
C_ACCENT   = RGBColor(0x2E, 0x74, 0xB5)

# ─── helpers ─────────────────────────────────────────────────────────────────
def set_cell_bg(cell, rgb: RGBColor):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  f'{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}')
    tcPr.append(shd)

def cell_para(cell, text, bold=False, color=C_BLACK, size=10, align=WD_ALIGN_PARAGRAPH.LEFT, italic=False):
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    run.bold   = bold
    run.italic = italic
    run.font.size  = Pt(size)
    run.font.color.rgb = color

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f'Heading {level}'
    run = p.add_run(text)
    run.font.color.rgb = C_TITLE
    return p

def add_body(doc, text, size=10):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    return p

def add_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run('ℹ  ' + text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = C_ACCENT
    return p

def build_table(doc, headers, rows, col_widths=None, zebra=True):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'

    # header row
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        set_cell_bg(hdr.cells[i], C_HDR_BG)
        cell_para(hdr.cells[i], h, bold=True, color=C_WHITE, size=10)

    # data rows
    for ri, row in enumerate(rows):
        bg = C_ROW_ALT if (zebra and ri % 2 == 1) else C_WHITE
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            set_cell_bg(cell, bg)
            # colour-code sel descriptions
            bold = False
            col  = C_BLACK
            if isinstance(val, tuple):
                val, col, bold = val
            cell_para(cell, str(val), bold=bold, color=col, size=9.5)

    # column widths
    if col_widths:
        for ri in range(len(table.rows)):
            for ci, w in enumerate(col_widths):
                table.rows[ri].cells[ci].width = Inches(w)

    doc.add_paragraph()   # spacer
    return table

# ─── document ────────────────────────────────────────────────────────────────
doc = Document()

# page margins
for section in doc.sections:
    section.top_margin    = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin   = Cm(2.2)
    section.right_margin  = Cm(2.2)

# ── title page ───────────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('PAL / NTSC Video IP Core')
r.bold = True; r.font.size = Pt(22); r.font.color.rgb = C_TITLE

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Configuration Provisions Reference')
r.bold = True; r.font.size = Pt(16); r.font.color.rgb = C_ACCENT

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Top-level entity: pal_tv_bram_top  /  video_bram_top')
r.font.size = Pt(10); r.italic = True; r.font.color.rgb = RGBColor(0x44,0x44,0x44)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════
# 1. GENERICS
# ═══════════════════════════════════════════════════════════════════════════
add_heading(doc, '1  Generics  (Compile-time / Vivado Block Properties)', 1)
add_body(doc,
    'Generics are set once at synthesis time — in Vivado via the IP Customisation GUI '
    'or in the generic map of an instantiation. They cannot be changed at runtime.')
doc.add_paragraph()

generics = [
    # (Generic, Default, Description)
    ('H_FRONT',     '16',     'Horizontal front porch — clock cycles before H-sync'),
    ('H_SYNC_W',    '47',     'H-sync pulse width in clock cycles'),
    ('H_BACK',      '57',     'H back porch — clock cycles after H-sync, before active'),
    ('H_ACTIVE',    '520',    'Active pixels per line (PAL standard)'),
    ('H_TOTAL',     '640',    'Total clocks per line  (H_FRONT + H_SYNC_W + H_BACK + H_ACTIVE)'),
    ('V_TOTAL',     '625',    'Total lines per frame  (625 = PAL,  525 = NTSC)'),
    ('LEVEL_SYNC',  '"0000"', '4-bit DAC code for sync tip  (0 V on R-2R ladder)'),
    ('LEVEL_BLANK', '"0100"', '4-bit DAC code for blanking floor  (~0.3 V)'),
    ('LEVEL_WHITE', '"1111"', '4-bit DAC code for peak white  (1 V)'),
    ('CLK_MHZ',     '40',     'Input clock frequency in MHz — used for ball animation timing'),
    ('STRIPE_W',    '4',      'Stripe width in pixels used by sel=0 and sel=1 patterns'),
    ('BRAM_DEPTH',  '299520', 'BRAM address space (520 × 576 = 299,520). Use 524288 for power-of-2 BRAM'),
]

build_table(doc,
    headers=['Generic', 'Default', 'Description'],
    rows=generics,
    col_widths=[1.3, 1.0, 4.5])

add_note(doc, 'For NTSC set V_TOTAL=525, H_ACTIVE=640, H_TOTAL=858 (if using 13.5 MHz pixel clock).')

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════
# 2. RUNTIME PORTS
# ═══════════════════════════════════════════════════════════════════════════
add_heading(doc, '2  Runtime Input Ports  (Changeable While Running)', 1)

# 2.1 Pattern Select
add_heading(doc, '2.1  Pattern Select', 2)
build_table(doc,
    headers=['Port', 'Width', 'Default', 'Description'],
    rows=[
        ('sel',        '8-bit', '—',      'Pattern selector (see Section 3 for all values)'),
        ('brightness', '4-bit', '"1111"', 'White level DAC code override'),
        ('black_lvl',  '4-bit', '"0100"', 'Black / blanking floor DAC code override'),
    ],
    col_widths=[1.3, 0.8, 1.0, 4.1])

# 2.2 BRAM Write Interface
add_heading(doc, '2.2  BRAM Write Interface  (sel = 4)', 2)
add_body(doc,
    'Write pixel data directly into the internal BRAM at any time. '
    'Address is linear (row-major). The display read-pointer applies interlace correction automatically.')
doc.add_paragraph()
build_table(doc,
    headers=['Port', 'Width', 'Description'],
    rows=[
        ('bram_wr_en',   '1-bit',  'Write enable — pulse high for one clock to write one pixel'),
        ('bram_wr_addr', '19-bit', 'Target BRAM address  (0 … bram_len−1)'),
        ('bram_wr_data', '1-bit',  'Pixel value  (1 = white,  0 = black)'),
        ('bram_len',     '19-bit', 'Total valid pixels loaded — address wraps here (default = all 1s)'),
    ],
    col_widths=[1.5, 0.9, 4.8])

add_note(doc, 'bram_wr_en is edge-sensitive; hold address and data stable for at least one clock.')

doc.add_paragraph()

# 2.3 Bouncing Ball
add_heading(doc, '2.3  Bouncing Ball Control  (sel = 5)', 2)
add_body(doc,
    'All four ports are latched on the rising edge of rst. '
    'Assert rst for ≥ 3 clock cycles after changing any ball parameter to reload the values.')
doc.add_paragraph()
build_table(doc,
    headers=['Port', 'Width', 'Default', 'Description'],
    rows=[
        ('ball_w_i',   '10-bit', '60  (px)',   'Ball width in pixels  (1 – 519)'),
        ('ball_h_i',   '10-bit', '50  (rows)', 'Ball height in active rows  (1 – 287)'),
        ('ball_spd_h', '4-bit',  '3   (px/f)', 'Horizontal speed in pixels per frame  (1 – 15)'),
        ('ball_spd_v', '4-bit',  '2   (row/f)','Vertical speed in rows per frame  (1 – 15)'),
    ],
    col_widths=[1.3, 0.9, 1.3, 3.7])

doc.add_paragraph()

# 2.4 Crosshatch
add_heading(doc, '2.4  Crosshatch Grid Control  (sel = 8)', 2)
add_body(doc,
    'Grid line spacing is set independently for horizontal and vertical. '
    'A grid line is always drawn at pixel 0 / row 0 of the active area; '
    'the spacing defines the period of the repeating grid.')
doc.add_paragraph()
build_table(doc,
    headers=['Port', 'Width', 'Default', 'Description'],
    rows=[
        ('cross_h_i', '10-bit', '52  (px)',   'Horizontal grid spacing — pixels between vertical grid lines'),
        ('cross_v_i', '10-bit', '58  (rows)', 'Vertical grid spacing — active rows between horizontal grid lines'),
    ],
    col_widths=[1.3, 0.9, 1.3, 3.7])

add_note(doc, 'Both values are clamped to a minimum of 1 internally to avoid counter wrap-around.')

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════
# 3. SEL VALUES
# ═══════════════════════════════════════════════════════════════════════════
add_heading(doc, '3  Pattern Selector (sel) Values', 1)
add_body(doc,
    'The 8-bit sel port chooses which pixel generator drives dac_out during the active picture area. '
    'Sync, blanking and back-porch levels are unaffected by sel.')
doc.add_paragraph()

SEL_ROWS = [
    ('0x00', '0', 'Vertical stripes',     'White stripes of width STRIPE_W on black background'),
    ('0x01', '1', 'Horizontal stripes',   'White stripes of height STRIPE_W on black background'),
    ('0x02', '2', 'Horizontal gradient',  'Luminance ramps left → right  (dark to bright)'),
    ('0x03', '3', 'Vertical gradient',    'Luminance ramps top → bottom  (dark to bright)'),
    ('0x04', '4', 'BRAM image',           'Displays 1-bit image written via BRAM write port (interlaced)'),
    ('0x05', '5', 'Bouncing ball',        'Solid white rectangle that bounces around screen; size/speed configurable at runtime'),
    ('0x06', '6', 'Full white',           'Every active pixel = LEVEL_WHITE  (DAC output calibration)'),
    ('0x07', '7', 'Full black',           'Every active pixel = LEVEL_BLANK  (DAC output calibration)'),
    ('0x08', '8', 'Crosshatch grid',      'Regular grid of white lines on black; H and V pitch set via cross_h_i / cross_v_i'),
    ('0x09', '9', 'Centre cross',         'Single white vertical line at H centre + single white horizontal line at V centre'),
]

table = doc.add_table(rows=1 + len(SEL_ROWS), cols=4)
table.style = 'Table Grid'

# header
hrow = table.rows[0]
for i, h in enumerate(['sel hex', 'sel dec', 'Pattern Name', 'Description']):
    set_cell_bg(hrow.cells[i], C_HDR_BG)
    cell_para(hrow.cells[i], h, bold=True, color=C_WHITE, size=10)

BRAM_COL  = RGBColor(0x1F, 0x49, 0x7D)
BALL_COL  = C_GREEN
CAL_COL   = C_ORANGE
CROSS_COL = RGBColor(0x70, 0x30, 0xA0)
GRAD_COL  = RGBColor(0x00, 0x60, 0x60)
STR_COL   = RGBColor(0x60, 0x30, 0x00)

ACCENT_COLORS = [STR_COL, STR_COL, GRAD_COL, GRAD_COL,
                 BRAM_COL, BALL_COL, CAL_COL, CAL_COL,
                 CROSS_COL, CROSS_COL]

for ri, (hex_v, dec_v, name, desc) in enumerate(SEL_ROWS):
    bg  = C_ROW_ALT if ri % 2 == 1 else C_WHITE
    ac  = ACCENT_COLORS[ri]
    row = table.rows[ri + 1]
    set_cell_bg(row.cells[0], bg); cell_para(row.cells[0], hex_v, bold=True, color=ac,      size=9.5)
    set_cell_bg(row.cells[1], bg); cell_para(row.cells[1], dec_v, bold=False, color=C_BLACK, size=9.5)
    set_cell_bg(row.cells[2], bg); cell_para(row.cells[2], name,  bold=True,  color=ac,      size=9.5)
    set_cell_bg(row.cells[3], bg); cell_para(row.cells[3], desc,  bold=False, color=C_BLACK, size=9.5)

for ri in range(len(table.rows)):
    for ci, w in enumerate([0.9, 0.8, 1.5, 4.0]):
        table.rows[ri].cells[ci].width = Inches(w)

doc.add_paragraph()
add_note(doc, 'sel values 0x0A and above are undefined — dac_out defaults to LEVEL_BLANK.')

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════════════════
# 4. OUTPUT PORTS
# ═══════════════════════════════════════════════════════════════════════════
add_heading(doc, '4  Output Ports', 1)
add_body(doc,
    'All outputs are synchronous to clk. Connect dac_out to a 4-bit R-2R resistor ladder '
    'to produce an analogue video signal. The sync outputs can drive standard video encoders '
    'or be combined externally.')
doc.add_paragraph()

build_table(doc,
    headers=['Port', 'Width', 'Description'],
    rows=[
        ('dac_out',      '4-bit', '4-bit DAC value — connect to R-2R resistor ladder to produce analogue video'),
        ('csync_o',      '1-bit', 'Composite sync (H and V merged) — high during sync tip'),
        ('line_sync_o',  '1-bit', 'H-sync only — high during horizontal sync, never during V-sync'),
        ('frame_sync_o', '1-bit', 'V-sync pulse — high during the broad-sync region of the V interval'),
        ('field_o',      '1-bit', 'Field indicator  (0 = Field 1 / odd lines,  1 = Field 2 / even lines)'),
        ('active_o',     '1-bit', 'High during the active picture area (H and V both in active window)'),
        ('blank_o',      '1-bit', 'Composite blanking — high outside active area, but NOT during sync tip'),
    ],
    col_widths=[1.4, 0.8, 5.0])

doc.add_paragraph()

# ═══════════════════════════════════════════════════════════════════════════
# 5. QUICK-REFERENCE SUMMARY
# ═══════════════════════════════════════════════════════════════════════════
add_heading(doc, '5  Quick-Reference Summary', 1)

summary = [
    ('Timing',          'H_FRONT, H_SYNC_W, H_BACK, H_ACTIVE, H_TOTAL, V_TOTAL',         'Generics — set at synthesis'),
    ('DAC levels',      'LEVEL_SYNC, LEVEL_BLANK, LEVEL_WHITE',                           'Generics — set at synthesis'),
    ('Clock',           'CLK_MHZ',                                                         'Generic — must match actual clock'),
    ('Pattern',         'sel  (8-bit)',                                                    'Runtime — change any time'),
    ('Brightness',      'brightness, black_lvl  (4-bit each)',                            'Runtime — change any time'),
    ('BRAM image',      'bram_wr_en / bram_wr_addr / bram_wr_data / bram_len',           'Runtime — write pixel-by-pixel'),
    ('Ball size/speed', 'ball_w_i, ball_h_i  (10-bit)   ball_spd_h, ball_spd_v  (4-bit)','Runtime — apply rst after change'),
    ('Crosshatch pitch','cross_h_i, cross_v_i  (10-bit each)',                           'Runtime — change any time'),
]

build_table(doc,
    headers=['Category', 'Ports / Generics', 'When to Configure'],
    rows=summary,
    col_widths=[1.5, 3.8, 1.9])

# ─── save ────────────────────────────────────────────────────────────────────
doc.save(OUT)
size = __import__('os').path.getsize(OUT)
print(f'Saved: {OUT}  ({size:,} bytes)')
